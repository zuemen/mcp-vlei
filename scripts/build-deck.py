"""Build the talk deck; its speaker notes are read from the script in `docs/DEMO.md`.

The slides and the spoken script have to stay in step, so the deck is generated rather than drawn,
and the notes are not copied into this file: each slide's notes are the block quotes under its
`### N — Title` heading in `docs/DEMO.md`, read at build time. A slide whose heading is missing, or
whose number in the script does not match its position in the deck, stops the build. Change the talk
there, run this, and the deck follows.

The output is native PowerPoint — real text boxes, real tables, editable on the machine that will
be plugged into the projector. Nothing on a slide is a picture of text. The pictures there are are
evidence: screenshots of the running console, taken by `scripts/capture-console.py` rather than
drawn, and a QR code for the repository.

    python scripts/build-deck.py                                   # -> docs/slides/mcp-vlei.pptx
    python scripts/build-deck.py --render                          # ... and contact sheets
    python scripts/build-deck.py --video docs/slides/demo-full.mp4 # embed the recording

`--video` embeds the recording on the Demo slide in the 16:9 frame the still occupies until then,
set to play full screen; with PowerPoint installed it is also set to start on entry and not to
rewind, and the result is read back to confirm it. Without a video the frame shows scene 3 — the
fallback the talk falls back to if a projector will not play video.

`--render` is not decoration. Every dimension in this file is arithmetic, and arithmetic produces
footers that collide with page numbers and beige panels with three empty inches under them — both
of which were in the first build and neither of which is visible from the code. It needs
LibreOffice on PATH (or at its usual Windows location) and PyMuPDF.
"""

from __future__ import annotations

import io
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "slides" / "mcp-vlei.pptx"
SCRIPT = ROOT / "docs" / "DEMO.md"
SHOTS = ROOT / "docs" / "slides" / "shots"
REPO_URL = "https://github.com/zuemen/mcp-vlei"

# The console's palette, so the deck and the thing on screen behind it are the same object.
PAPER = RGBColor(0xF7, 0xF5, 0xEF)
INK = RGBColor(0x1E, 0x1B, 0x16)
MUTED = RGBColor(0x6E, 0x66, 0x58)
BLUE = RGBColor(0x2E, 0x5D, 0x89)
RED = RGBColor(0xC0, 0x39, 0x2B)
RULE = RGBColor(0xDC, 0xD6, 0xC8)

SERIF = "Georgia"
SANS = "Segoe UI"
MONO = "Consolas"
#: Neither Georgia nor Segoe UI carries CJK glyphs; without an East Asian font on the run, the
#: renderer substitutes whatever it finds, and on a borrowed machine that can be boxes.
CJK = "Microsoft JhengHei"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.95)
BODY_W = W - 2 * MARGIN

# Everything below the rule and above the footer. Content is centred in it rather than hung from
# the top: a four-line slide otherwise sits in the upper third with the lower half empty, which on
# a projector reads as a slide that is missing something.
BODY_TOP = Inches(2.16)
BODY_BOTTOM = Inches(6.34)
BODY_H = BODY_BOTTOM - BODY_TOP

#: A picture beside a table or a panel: this far from it.
GAP = Inches(0.4)

# The title slide. The title is set as two lines on purpose — at 44pt it does not fit on one, and
# letting it wrap is how the subtitle came to be printed over its second line. Fixed lines, fixed
# leading, and the subtitle placed from their arithmetic.
TITLE_LINES = ("Organizational Identity", "for MCP Agents")
TITLE_SIZE = 44
TITLE_LEADING = Pt(52)
TITLE_TOP = Inches(1.9)
TITLE_BOTTOM = TITLE_TOP + TITLE_LEADING * len(TITLE_LINES)
SUBTITLE_TOP = TITLE_BOTTOM + Inches(0.3)  # at least 0.25" below the title's last line

# The demo frame: 16:9, about four fifths of the content width, the recording being the slide.
VIDEO_W = Inches(9.2)
VIDEO_H = Emu(int(VIDEO_W * 9 / 16))
VIDEO_LEFT = Emu(int((W - VIDEO_W) / 2))
VIDEO_TOP = Inches(0.55)
DEMO_STILL = SHOTS / "scene-3.png"


# ------------------------------------------------------------------------------------------- #
# Primitives
# ------------------------------------------------------------------------------------------- #

def _text(slide, left, top, width, height, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.paragraphs[0].alignment = align
    return frame


def _has_cjk(text: str) -> bool:
    return any("\u2e80" <= ch <= "\u9fff" or "\uf900" <= ch <= "\ufaff" or "\uff00" <= ch <= "\uffef"
               for ch in text)


def _run(paragraph, text, *, size, color=INK, font=SANS, bold=False, italic=False, spacing=None):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    run.font.bold = bold
    run.font.italic = italic
    if spacing is not None:
        # python-pptx has no character-spacing API; set it on the run's rPr directly.
        run.font._rPr.set("spc", str(int(spacing * 100)))
    if _has_cjk(text):
        rpr = run.font._rPr
        latin = rpr.find(qn("a:latin"))
        ea = rpr.makeelement(qn("a:ea"), {"typeface": CJK})
        if latin is not None:
            latin.addnext(ea)
        else:
            rpr.append(ea)
    return run


def _rule(slide, top, width=None, color=RULE, height=Emu(9525), left=None):
    line = slide.shapes.add_shape(1, left if left is not None else MARGIN, top,
                                  width or BODY_W, height)  # 1 = rectangle
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()
    line.shadow.inherit = False
    return line


def _picture(slide, path: Path, left, top, width):
    """A screenshot, with a hairline so a white UI does not dissolve into the paper colour."""
    if not path.is_file():
        raise SystemExit(f"{path} is missing; run `python scripts/capture-console.py` first")
    picture = slide.shapes.add_picture(str(path), left, top, width=width)
    picture.line.color.rgb = RULE
    picture.line.width = Pt(0.75)
    return picture


def _qr(slide, left, top, size):
    """The repository's URL as a QR code — the one thing on a slide meant for a phone."""
    import segno

    buffer = io.BytesIO()
    segno.make(REPO_URL, error="m").save(buffer, kind="png", scale=12, border=0,
                                         dark="#1E1B16", light="#F7F5EF")
    buffer.seek(0)
    return slide.shapes.add_picture(buffer, left, top, width=size, height=size)


def _slide(deck, notes=""):
    slide = deck.slides.add_slide(deck.slide_layouts[6])  # blank
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER
    if notes:
        slide.notes_slide.notes_text_frame.text = notes.strip()
    return slide


def _heading(slide, kicker, title):
    """The running header: a small tracked-out kicker over a serif title."""
    frame = _text(slide, MARGIN, Inches(0.62), BODY_W, Inches(0.3))
    _run(frame.paragraphs[0], kicker.upper(), size=11, color=BLUE, bold=True, spacing=1.6)

    frame = _text(slide, MARGIN, Inches(1.0), BODY_W, Inches(0.9))
    _run(frame.paragraphs[0], title, size=32, font=SERIF)
    _rule(slide, Inches(1.86))


def _footer(slide, text, *, color=MUTED):
    # Stop short of the page number: a two-line footer used to run straight through it.
    frame = _text(slide, MARGIN, Inches(6.5), BODY_W - Inches(1.5), Inches(0.5))
    _run(frame.paragraphs[0], text, size=13, color=color, italic=True, font=SERIF)


def _page(slide, n, total):
    frame = _text(slide, W - MARGIN - Inches(1.2), Inches(6.54), Inches(1.2), Inches(0.3),
                  align=PP_ALIGN.RIGHT)
    _run(frame.paragraphs[0], f"{n} / {total}", size=10, color=MUTED, spacing=0.8)


#: "No Style, No Grid" — the default style draws white borders, invisible on this paper.
_PLAIN_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _rows_ruled(table):
    """A hairline under every row and nothing else: rows separated, columns by alignment alone."""
    style = table._tbl.tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = table._tbl.tblPr.makeelement(qn("a:tableStyleId"), {})
        table._tbl.tblPr.append(style)
    style.text = _PLAIN_TABLE
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                for old in tc_pr.findall(qn(tag)):
                    tc_pr.remove(old)
            for position, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                line = tc_pr.makeelement(qn(tag), {"w": "9525" if tag == "a:lnB" else "0"})
                if tag == "a:lnB":
                    fill = line.makeelement(qn("a:solidFill"), {})
                    color = fill.makeelement(qn("a:srgbClr"), {"val": "DCD6C8"})
                    fill.append(color)
                    line.append(fill)
                else:
                    line.append(line.makeelement(qn("a:noFill"), {}))
                tc_pr.insert(position, line)


def _beside(spec):
    """Split the body when a slide carries a picture: (content width, picture left, picture width)."""
    picture = spec.get("picture")
    if not picture:
        return BODY_W, None, None
    width = Inches(picture["width"])
    return BODY_W - width - GAP, MARGIN + BODY_W - width, width


def _place_picture(slide, spec, left, width):
    picture = spec["picture"]
    shape = _picture(slide, SHOTS / picture["file"], left, BODY_TOP, width)
    room = BODY_H - (Inches(0.38) if picture.get("caption") else 0)
    if shape.height > room:  # a tall crop is fitted to the body, and centred in its column
        shape.width = Emu(int(shape.width * room / shape.height))
        shape.height = room
        shape.left = left + (width - shape.width) // 2
    shape.top = BODY_TOP + (room - shape.height) // 2
    if picture.get("caption"):
        frame = _text(slide, shape.left, shape.top + shape.height + Inches(0.1), shape.width,
                      Inches(0.3))
        _run(frame.paragraphs[0], picture["caption"], size=10, color=MUTED)


# ------------------------------------------------------------------------------------------- #
# Slide kinds
# ------------------------------------------------------------------------------------------- #

def title_slide(deck, spec):
    slide = _slide(deck, spec["notes"])
    frame = _text(slide, MARGIN, TITLE_TOP, BODY_W, TITLE_BOTTOM - TITLE_TOP)
    for i, line in enumerate(TITLE_LINES):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.line_spacing = TITLE_LEADING
        _run(p, line, size=TITLE_SIZE, font=SERIF)

    frame = _text(slide, MARGIN, SUBTITLE_TOP, BODY_W, Inches(0.5))
    _run(frame.paragraphs[0], spec["subtitle"], size=19, color=BLUE)

    rule_top = SUBTITLE_TOP + Inches(0.78)
    _rule(slide, rule_top, width=Inches(2.2), color=BLUE, height=Inches(0.03))

    frame = _text(slide, MARGIN, rule_top + Inches(0.42), BODY_W - Inches(1.6), Inches(1.0))
    for i, line in enumerate(spec["meta"]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(6)
        _run(p, line, size=14, color=MUTED, font=MONO if line.startswith("github") else SANS)

    size = Inches(1.15)
    _qr(slide, W - MARGIN - size, rule_top + Inches(0.42), size)


def bullets_slide(deck, spec):
    slide = _slide(deck, spec["notes"])
    _heading(slide, spec["kicker"], spec["title"])

    frame = _text(slide, MARGIN, BODY_TOP, BODY_W, BODY_H, anchor=MSO_ANCHOR.MIDDLE)
    for i, item in enumerate(spec["bullets"]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(spec.get("gap", 16))
        if isinstance(item, tuple):
            lead, rest = item
            _run(p, lead, size=spec.get("size", 20), color=BLUE, bold=True)
            _run(p, rest, size=spec.get("size", 20))
        else:
            _run(p, item, size=spec.get("size", 20))
    if spec.get("footer"):
        _footer(slide, spec["footer"])


def table_slide(deck, spec):
    slide = _slide(deck, spec["notes"])
    _heading(slide, spec["kicker"], spec["title"])
    table_w, picture_left, picture_w = _beside(spec)

    rows, cols = len(spec["rows"]) + 1, len(spec["head"])
    height = Inches(spec.get("row_height", 0.46)) * rows
    top = BODY_TOP + (BODY_H - height) // 2
    shape = slide.shapes.add_table(rows, cols, MARGIN, top, table_w, height)
    table = shape.table
    table.first_row = True
    for c, width in zip(table.columns, spec["widths"]):
        c.width = Emu(int(table_w * width))

    for c, label in enumerate(spec["head"]):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PAPER
        frame = cell.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.1)
        # Latin and CJK in one header are set as two runs: tracking and capitals for the Latin,
        # neither for the CJK, which tracking pulls apart.
        cut = next((i for i, ch in enumerate(label) if _has_cjk(ch)), len(label))
        latin, cjk = label[:cut].rstrip(), label[cut:]
        _run(frame.paragraphs[0], latin.upper(), size=11, color=BLUE, bold=True, spacing=1.2)
        if cjk:
            _run(frame.paragraphs[0], "  " + cjk, size=11, color=BLUE, bold=True)

    for r, row in enumerate(spec["rows"], start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PAPER
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame = cell.text_frame
            frame.word_wrap = True
            frame.margin_left = frame.margin_right = Inches(0.1)
            mono = value.startswith("`") and value.endswith("`")  # the whole cell, or none of it
            negative = value.startswith("— ")  # "—" alone means "nothing today", not a refusal
            _run(frame.paragraphs[0], value.replace("`", ""),
                 size=spec.get("size", 14),
                 color=RED if negative else INK,
                 font=MONO if mono else SANS)
    _rows_ruled(table)
    if spec.get("row_height"):
        table.rows[0].height = Inches(0.46)  # the header stays the height of every other header
    if picture_left is not None:
        _place_picture(slide, spec, picture_left, picture_w)
    if spec.get("footer"):
        _footer(slide, spec["footer"])


def mono_slide(deck, spec):
    slide = _slide(deck, spec["notes"])
    _heading(slide, spec["kicker"], spec["title"])
    panel_w, picture_left, picture_w = _beside(spec)
    if spec.get("reserve"):
        panel_w -= Inches(spec["reserve"])

    size = spec.get("size", 15)
    pad = Inches(0.32)
    # A panel with three inches of empty beige under the last line looks like a rendering failure.
    height = pad * 2 + Pt(size * 1.3 + 3) * len(spec["lines"])
    top = BODY_TOP + (BODY_H - height) // 2

    panel = slide.shapes.add_shape(1, MARGIN, top, panel_w, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(0xEF, 0xEB, 0xE1)
    panel.line.fill.background()
    panel.shadow.inherit = False

    frame = _text(slide, MARGIN + pad, top + pad, panel_w - pad * 2, height - pad * 2)
    for i, line in enumerate(spec["lines"]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(3)
        # A leading marker colours the key rather than the whole line.
        if line.startswith("*"):
            _run(p, line[1:], size=size, color=BLUE, font=MONO, bold=True)
        else:
            _run(p, line, size=size, color=INK, font=MONO)
    if picture_left is not None:
        _place_picture(slide, spec, picture_left, picture_w)
    if spec.get("qr"):
        qr = Inches(spec["qr"])
        left = W - MARGIN - qr
        caption_gap, caption_h = Inches(0.28), Inches(0.26)
        qr_top = top + (height - qr - caption_gap - caption_h) // 2
        _qr(slide, left, qr_top, qr)
        # Centred on the code, with a quiet zone of about four modules above the caption.
        frame = _text(slide, left - Inches(0.4), qr_top + qr + caption_gap, qr + Inches(0.8),
                      caption_h, align=PP_ALIGN.CENTER)
        _run(frame.paragraphs[0], REPO_URL.removeprefix("https://"), size=9.5, color=MUTED,
             font=MONO)
    if spec.get("footer"):
        _footer(slide, spec["footer"], color=INK)


def demo_slide(deck, spec):
    """The recording is the slide. Until it exists, the still of scene 3 holds its frame."""
    slide = _slide(deck, spec["notes"])
    video = spec.get("video")
    if video:
        movie = slide.shapes.add_movie(str(video), VIDEO_LEFT, VIDEO_TOP, VIDEO_W, VIDEO_H,
                                       poster_frame_image=str(DEMO_STILL), mime_type="video/mp4")
        # Full screen is a property of the timing node python-pptx writes for the movie.
        for node in slide.element.iter(qn("p:video")):
            node.set("fullScrn", "1")
        movie.name = "demo-full"
    else:
        _picture(slide, DEMO_STILL, VIDEO_LEFT, VIDEO_TOP, VIDEO_W)

    frame = _text(slide, VIDEO_LEFT, VIDEO_TOP + VIDEO_H + Inches(0.36), VIDEO_W - Inches(1.3),
                  Inches(0.7))
    for i, line in enumerate(spec["lines"]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(2)
        _run(p, line, size=15, font=SERIF, color=RED if line.startswith("The root") else INK)


def statement_slide(deck, spec):
    """One sentence, alone."""
    slide = _slide(deck, spec["notes"])
    frame = _text(slide, MARGIN, Inches(2.6), BODY_W, Inches(0.4))
    _run(frame.paragraphs[0], spec["kicker"].upper(), size=11, color=BLUE, bold=True, spacing=1.6)

    frame = _text(slide, MARGIN, Inches(3.2), BODY_W, Inches(2.2))
    for i, line in enumerate(spec["lines"]):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.space_after = Pt(14)
        _run(p, line, size=spec.get("size", 30), font=SERIF,
             color=RED if line.startswith("The root") else INK)
    if spec.get("footer"):
        _footer(slide, spec["footer"])


KINDS = {"title": title_slide, "bullets": bullets_slide, "table": table_slide,
         "mono": mono_slide, "statement": statement_slide, "demo": demo_slide}


# ------------------------------------------------------------------------------------------- #
# The notes: read from docs/DEMO.md
# ------------------------------------------------------------------------------------------- #

def load_notes(script: Path = SCRIPT) -> dict[int, tuple[str, str]]:
    """``{slide number: (heading, notes)}`` from the ``# Slides`` section of the talk script.

    Notes are the block quotes under each heading, with the quotation marks and Markdown taken off.
    A stage direction on its own line — ``*(Play demo-full.mp4, 4:05.)*`` — is kept, bracketed, so
    the presenter view says when to press play.
    """
    text = script.read_text(encoding="utf-8")
    start = text.index("\n# Slides")
    end = text.index("\n# Anticipated questions")
    section = text[start:end]

    notes: dict[int, tuple[str, str]] = {}
    for match in re.finditer(r"^### (\d+) — (.+?)\n(.*?)(?=^### |\Z)", section, re.S | re.M):
        number, heading, body = int(match.group(1)), match.group(2).strip(), match.group(3)
        paragraphs: list[str] = []
        current: list[str] = []
        direction: list[str] = []
        for line in body.splitlines():
            # A stage direction may run over several lines: *( … )*
            if direction or line.strip().startswith("*("):
                direction.append(line.strip())
                if line.strip().endswith(")*"):
                    joined = " ".join(direction)
                    paragraphs.append(f"[{joined[2:-2].strip()}]")
                    direction = []
                continue
            if line.startswith(">"):
                content = line[1:].strip()
                if content:
                    current.append(content)
                elif current:
                    paragraphs.append(" ".join(current))
                    current = []
            else:
                if current:
                    paragraphs.append(" ".join(current))
                    current = []
        if current:
            paragraphs.append(" ".join(current))
        cleaned = []
        for paragraph in paragraphs:
            paragraph = paragraph.replace("`", "")
            paragraph = re.sub(r"\*([^*]+)\*", r"\1", paragraph)
            if paragraph.startswith('"'):
                paragraph = paragraph[1:]
            if paragraph.endswith('"'):
                paragraph = paragraph[:-1]
            cleaned.append(paragraph)
        notes[number] = (heading, "\n\n".join(cleaned))
    return notes


# ------------------------------------------------------------------------------------------- #
# The deck. `heading` is the slide's `### N — heading` in docs/DEMO.md; it defaults to `title`.
# ------------------------------------------------------------------------------------------- #

SLIDES = [
    {"kind": "title", "heading": "Title",
     "subtitle": "vLEI × Model Context Protocol",
     "meta": ["An additive extension: org.gleif.vlei/identity",
              "github.com/zuemen/mcp-vlei · Apache 2.0"]},

    {"kind": "table", "kicker": "The stack today", "title": "What MCP verifies",
     "head": ["Layer", "What it proves", "What it does not"],
     "widths": [0.22, 0.39, 0.39],
     "rows": [
         ["TLS", "Control of a domain name", "— Which legal entity"],
         ["OAuth `iss`", "Which authorization server", "— Which legal entity"],
         ["OAuth `client_id`", "Control of a metadata URL", "— Which legal entity"],
         ["OAuth `sub`", "Which human user", "— Which organization they bind"],
         ["`clientInfo`", "Nothing. It is self-asserted", "— MUST NOT be used for security"],
     ],
     "footer": "Domain control, domain control, domain control, human user."},

    {"kind": "table", "kicker": "Measured, not asserted", "title": "The measurement",
     # The names are what the vendor server received, from examples/impersonation/spoof_client.py.
     # Omitting clientInfo does not send nothing: the SDK supplies `mcp`.
     "head": ["Run", "Name received", "Approved"],
     "widths": [0.30, 0.42, 0.28],
     "rows": [["Honest", "`zuemen-script`", "1 hour"],
              ["Spoofed", "`Claude Desktop`", "50 hours"],
              ["Omitted", "`mcp`", "1 hour"]],
     "size": 15,
     "picture": {"file": "scene-0-verification.png", "width": 4.2,
                 "caption": "Scene 0 · nothing to check, and 50 hours granted"},
     "footer": "One binary, one variable. The server's policy deliberately violates the "
               "specification — that is the experiment."},

    {"kind": "mono", "kicker": "The protocol itself",
     "title": "The agent is absent from the protocol",
     # A leading * marks a line and is removed, so marked lines carry one more space.
     "lines": ["schema.ts · revision 2026-07-28 · 3,197 lines", "",
               "*   agent        0 occurrences",
               "*   principal    0 occurrences",
               "*   delegation   0 occurrences",
               "*   mandate      0 occurrences", "",
               "   host         modelled",
               "   client       modelled",
               "   server       modelled"],
     "footer": "It cannot be named, delegated to, constrained, or revoked — there is nothing "
               "there to name."},

    {"kind": "bullets", "kicker": "Read it fairly",
     "title": "This is a boundary, not a defect",
     "bullets": [
         ('"MCP\'s trust model assumes a human in the loop"',
          " — the Tools chapter's warning box"),
         ("NSA, May 2026", " — MCP does not define how a session maps to a verifiable identity"),
         ("Security IG proposals", " — root in domain, DNS, or registry"),
         ("Those answer ", "which deployment is this. None answers which legal entity is this."),
     ],
     "size": 19},

    {"kind": "bullets", "kicker": "Why this year", "title": "Why now",
     "bullets": ["Autonomous execution — no human at the point of action",
                 "Actions with legal effect — a filing, a payment, a submission",
                 "Across organizations — a counterparty that has never met the agent"],
     "gap": 22, "size": 22,
     "footer": "Any one alone is survivable. Together they are not."},

    {"kind": "bullets", "kicker": "Not a new identity system",
     "title": "GLEIF already solved the identity half",
     "bullets": [("Legal Entity (LE)", " — binds an identifier to an LEI"),
                 ("Engagement Context Role (ECR)", " — binds a person's identifier to a role "
                                                   "within that entity"),
                 ("Both chained ACDCs", " · both revocable · both verifiable offline"),
                 ("Anchored in KERI", " — key event logs, not a certificate authority")]},

    {"kind": "table", "kicker": "A deliberate choice", "title": "Why ECR and not OOR",
     "head": ["", "OOR", "ECR"],
     "widths": [0.24, 0.38, 0.38],
     "rows": [["Names", "A public office", "An engagement context"],
              ["Vocabulary", "Controlled", "Defined by the entity"],
              ["Validated by", "GLEIF", "The entity itself"]],
     "size": 16,
     "footer": "\"May file regulatory returns up to this amount\" is an engagement context, not a "
               "public office."},

    {"kind": "mono", "kicker": "The extension", "title": "What we added: the schema",
     "lines": ['capabilities.extensions["org.gleif.vlei/identity"]',
               "    presents / requires / acceptedRoots / ttlMs", "",
               "params._meta",
               "*    org.gleif.vlei/credential",
               "*    org.gleif.vlei/credentialSaid",
               "*    org.gleif.vlei/delegatedAid",
               "*    org.gleif.vlei/signature", "",
               "Tool._meta",
               "*    org.gleif.vlei/requires  { credential, role, scope }"],
     "size": 14,
     "footer": "Four keys in _meta. Nothing in core changes."},

    {"kind": "mono", "kicker": "The extension", "title": "What we added: verification",
     "lines": ["   1  credential_present   request",
               "   2  freshness            request",
               "   3  digest               request",
               "*   4  signature            signer's key log · witness",
               "*   5  delegation           holder's key log",
               "   6  chain                stream · issuers' logs",
               "   7  revocation           issuers' live logs",
               "   8  authority            credential"],
     "size": 13,
     "picture": {"file": "scene-1-verification.png", "width": 4.2,
                 "caption": "Scene 1 · the same checks, run on a live call"},
     "footer": "The key a request is checked against comes from the signer's own log — never "
               "from the request."},

    {"kind": "bullets", "kicker": "From building it",
     "title": "Three things the paper design did not predict",
     "bullets": [("1  ", "Presentation is the holder's step — so a counterparty's credential "
                         "must be verified locally"),
                 ("2  ", "Ask about the issuee, not the signer"),
                 ("3  ", "Protocol version negotiation decides whether extensions exist at all")],
     "gap": 24, "size": 21,
     "footer": "None of the three is derivable from the design. They are what a second "
               "implementer would otherwise pay for again."},

    {"kind": "bullets", "kicker": "Traceability", "title": "Every requirement, traced",
     "bullets": ["docs/CONFORMANCE.md — every normative statement, its code, its test",
                 "Three rows had nothing behind them. Writing the table is what found them.",
                 ("One defect every green test missed: ", "the signature was checked under a key "
                                                          "the request carried. Found, fixed, "
                                                          "recorded."),
                 "A section listing what is deliberately not claimed"],
     "gap": 18, "size": 20,
     "footer": "A conformance document that lists only successes is not evidence of anything."},

    {"kind": "demo", "heading": "Demo",
     "lines": ["Real KERI, real ACDC, real revocation.",
               "The root of trust is self-configured; in production it would be GLEIF's."]},

    {"kind": "table", "kicker": "Close to home",
     # Every cell is checked against MODA's own releases and code (docs/DEMO.md, slide 14,
     # "Sources"): the pilot's name, the credentials actually issued, the stack it runs on.
     "title": "Taiwan already runs this model",
     "head": ["", "Digital Identity Wallet 數位憑證皮夾", "This proposal"],
     "widths": [0.17, 0.43, 0.40],
     "rows": [["Holder", "People — and, since May 2026, a company's authorized representative",
               "Legal entities, and the agents acting for them"],
              ["Credentials", "Driving-licence verification card, degree certificates, "
                              "MOEA business certificate",
               "LE and ECR credentials"],
              ["Used for", "Parcel pickup at convenience stores, car-rental pilots",
               "Filing, verification, enquiry between agencies"],
              ["Built on", "Selective disclosure · OpenID4VC, SD-JWT VC, W3C VC",
               "Selective disclosure · KERI, ACDC — GLEIF's vLEI"]],
     "size": 14, "row_height": 0.62,
     "footer": "The trust model is already in use here. The holder it does not have yet is the "
               "agent acting for the entity."},

    {"kind": "bullets", "kicker": "Both already exist in institutions",
     "title": "Two ways to check an identity",
     "bullets": [("(a) Passive, from a public location", " — publish once, verify anywhere, no "
                                                          "per-check cost"),
                 ("(b) Attested confirmation (來函確認)", " — one institution confirms to "
                                                          "another, signed"),
                 ("The limit, stated: ", "relying on an attestation means relying on that "
                                         "institution's judgment — exactly as relying on a letter "
                                         "does today.")],
     "gap": 22, "size": 19,
     "footer": "The attesting institution must itself have been verified first."},

    {"kind": "bullets", "kicker": "Adoption", "title": "Six stages, one of which touches IT",
     "bullets": [("0  ", "Credential and role vocabulary — the business unit, not IT"),
                 ("1  ", "Publish at a well-known location"),
                 ("2  ", "Declare the requirement per tool"),
                 ("3  ", "Verify at the gateway ← the only IT change"),
                 ("4  ", "Attestations between institutions"),
                 ("5  ", "Audit records")],
     "gap": 12, "size": 19,
     "footer": "The filing service contains no verification code. It reads five headers the "
               "gateway sets."},

    {"kind": "table", "kicker": "Adoption", "title": "What an institution gets",
     "head": ["", "Today", "After stage 3"],
     "widths": [0.30, 0.35, 0.35],
     "rows": [["Who filed", "A session, a token", "A named entity and role"],
              ["Authority changed", "A ticket", "Revoke the credential"],
              ["New system built", "—", "None"],
              ["Existing sign-on", "Answers which user", "Unchanged — still required"],
              ["Audit record", "An IP and a timestamp", "LEI, role, signature"]],
     "size": 14,
     "footer": "vLEI answers which organization. Your existing sign-on still answers which user. "
               "A high-value action should require both."},

    {"kind": "bullets", "kicker": "Stated, not waited for", "title": "Limits",
     "bullets": ["Verifiable is not the same as trustworthy",
                 "LEIs are not issued to private individuals",
                 "There is a cost, and QVI coverage is still expanding",
                 "Our demo root is self-configured",
                 "One witness is asked — conflicting key logs are not yet detected",
                 "Revoking one agent's delegation is not implemented; revoking the credential is",
                 "Agent delegation conventions are not yet settled"],
     "gap": 9, "size": 18,
     "footer": "A valid credential proves an organization asserted a role. It does not prove the "
               "request is legitimate."},

    {"kind": "bullets", "kicker": "The ask", "title": "Three requests",
     "bullets": [("Government  ", "one small pilot: one filing or lookup procedure, stages 0 to "
                                  "3, no change to the department's systems"),
                 ("GLEIF  ", "an ECR role vocabulary for agent contexts · confirmation of the "
                             "delegation model against the EGF · test credentials"),
                 ("AAIF  ", "take this to the Security Interest Group and the ext-auth "
                            "discussion")],
     "gap": 26, "size": 20,
     "footer": "We chose a reading of the existing mechanisms, and would rather be corrected now "
               "than at deployment."},

    {"kind": "mono", "kicker": "One repository, Apache 2.0", "title": "Artifacts",
     "lines": ["spec/      specification v0.2 · types · wire examples · errors",
               "packages/  mcp-vlei — KEL + chain verification, 143 tests",
               "skills/    implementing-vlei (build) · vlei-identity (runtime)",
               "examples/  console · impersonation · regulator · skill-server",
               "deploy/    gateway configuration — zero-code-change adoption",
               "docs/      problem · adoption · conformance · upstream defect"],
     "size": 14, "reserve": 2.1, "qr": 1.6,
     "footer": "Real KERI, real ACDC, real revocation. The root of trust is self-configured; in "
               "production it would be GLEIF's."},
]


def apply_overrides(slides: list[dict], overrides: dict) -> list[dict]:
    """A variant of the deck without touching this file: ``{heading: {key: value}}`` replaces those
    keys of that slide, and ``"_drop": [heading, …]`` removes slides. Used for proposals, which live
    beside the talk rather than in it (docs/proposals/)."""
    drop = set(overrides.get("_drop", []))
    out = []
    for spec in slides:
        heading = spec.get("heading") or spec["title"]
        if heading in drop:
            continue
        changes = dict(overrides.get(heading, {}))
        if "rows" in changes:
            changes["rows"] = tuple(changes["rows"])
        if "bullets" in changes:  # JSON has no tuples; a [lead, rest] pair is a bullet with a lead
            changes["bullets"] = [tuple(b) if isinstance(b, list) else b for b in changes["bullets"]]
        out.append(dict(spec, **changes))
    return out


def build(path: Path = OUT, *, video: Path | None = None, script: Path = SCRIPT,
          overrides: dict | None = None) -> Path:
    notes = load_notes(script)
    deck = Presentation()
    deck.slide_width, deck.slide_height = W, H

    slides = apply_overrides(SLIDES, overrides) if overrides else SLIDES
    total = len(slides)
    for n, spec in enumerate(slides, start=1):
        heading = spec.get("heading") or spec["title"]
        if n not in notes or notes[n][0] != heading:
            found = notes.get(n, ("(nothing)",))[0]
            raise SystemExit(
                f"slide {n} is {heading!r}, but {script.name} has {found!r} as ### {n}. "
                "The script and the deck must be numbered the same way."
            )
        spec = dict(spec, notes=notes[n][1])
        if spec["kind"] == "demo" and video:
            spec["video"] = video
        KINDS[spec["kind"]](deck, spec)
        if n > 1:
            _page(deck.slides[-1], n, total)
    extra = sorted(set(notes) - set(range(1, total + 1)))
    if extra:
        raise SystemExit(f"{script.name} has notes for slides {extra}, which the deck does not have")

    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(path)
    return path


#: PowerPoint's msoAnimEffectMediaPlay and msoAnimTriggerWithPrevious.
MEDIA_PLAY, WITH_PREVIOUS = 83, 2


def set_playback(path: Path) -> dict[str, object]:
    """Start the recording on entry and do not rewind it — through PowerPoint, which writes the
    timing XML for this correctly — then read the settings back. Windows with PowerPoint only."""
    import win32com.client  # noqa: PLC0415

    app = win32com.client.Dispatch("PowerPoint.Application")
    presentation = app.Presentations.Open(str(path.resolve()), WithWindow=False)
    try:
        found: dict[str, object] = {}
        for slide in presentation.Slides:
            for shape in slide.Shapes:
                if shape.Name == "demo-full":
                    play = shape.AnimationSettings.PlaySettings
                    play.PlayOnEntry = True
                    play.RewindMovie = False
                    play.HideWhileNotPlaying = False
                    found = {"slide": slide.SlideIndex, "playOnEntry": bool(play.PlayOnEntry),
                             "rewind": bool(play.RewindMovie)}
            # PlayOnEntry alone leaves the play effect in the main sequence as a click effect: the
            # movie waited for a click, while the setting read back True. Make the effect start with
            # the slide (it is the slide's first effect, so "with previous" is "on entry").
            sequence = slide.TimeLine.MainSequence
            for index in range(1, sequence.Count + 1):
                effect = sequence.Item(index)
                if effect.Shape.Name == "demo-full" and effect.EffectType == MEDIA_PLAY:
                    effect.Timing.TriggerType = WITH_PREVIOUS
        presentation.Save()
    finally:
        presentation.Close()
    # PowerPoint rewrites the timing tree when it saves, and drops the full-screen flag with it; set
    # it again afterwards, then read both back from the file.
    deck = Presentation(str(path))
    for slide in deck.slides:
        for node in slide.element.iter(qn("p:video")):
            node.set("fullScrn", "1")
    deck.save(str(path))
    xml = "".join(s.element.xml for s in Presentation(str(path)).slides)
    found["fullScreen"] = 'fullScrn="1"' in xml
    # On entry means the play command's node is not a click effect; "playFrom" alone is also
    # present when the movie waits for a click.
    play_nodes = re.findall(r'<p:cTn [^>]*presetClass="mediacall"[^>]*nodeType="(\w+)"[^>]*>'
                            r'(?:(?!</p:cTn>).)*?playFrom', xml, re.S)
    found["startsOnEntry"] = bool(play_nodes) and all(n in ("withEffect", "afterEffect")
                                                      for n in play_nodes)
    return found


def render(deck_path: Path, out_dir: Path) -> list[Path]:
    """Rasterize into 2x2 contact sheets, and each slide on its own, to look at the result."""
    import shutil
    import subprocess

    import fitz

    soffice = shutil.which("soffice") or shutil.which("libreoffice") or \
        r"C:\Program Files\LibreOffice\program\soffice.exe"
    if not Path(soffice).exists():
        raise SystemExit(f"no LibreOffice at {soffice}; install it or render by hand")

    out_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir),
                    str(deck_path)], check=True, capture_output=True)

    pdf = fitz.open(out_dir / f"{deck_path.stem}.pdf")
    for n, page in enumerate(pdf, start=1):
        page.get_pixmap(dpi=110).save(out_dir / f"slide-{n:02d}.png")
    width, height = 1600, 900
    sheets = []
    for first in range(0, len(pdf), 4):
        sheet = fitz.open().new_page(width=width, height=height)
        for i, page in enumerate(range(first, min(first + 4, len(pdf)))):
            box = fitz.Rect((i % 2) * width / 2, (i // 2) * height / 2,
                            (i % 2 + 1) * width / 2, (i // 2 + 1) * height / 2)
            sheet.show_pdf_page(box, pdf, page)
        path = out_dir / f"sheet{first // 4}.png"
        sheet.get_pixmap(dpi=150).save(path)
        sheets.append(path)
    return sheets


def speaking_time(script: Path = SCRIPT, words_per_minute: int = 130) -> tuple[int, float]:
    """Words spoken from the notes — stage directions excluded — and minutes at that pace."""
    words = sum(len(re.findall(r"[A-Za-z0-9'’-]+", re.sub(r"\[[^\]]*\]", "", text)))
                for _, text in load_notes(script).values())
    return words, words / words_per_minute


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("out", nargs="?", type=Path, default=OUT)
    parser.add_argument("--render", action="store_true", help="also write PNG contact sheets")
    parser.add_argument("--video", type=Path, help="embed this recording on the Demo slide")
    parser.add_argument("--script", type=Path, default=SCRIPT,
                        help="read the notes from this script instead of docs/DEMO.md")
    parser.add_argument("--overrides", type=Path,
                        help="JSON of per-slide changes, for a variant of the deck (docs/proposals/)")
    args = parser.parse_args()

    import json

    overrides = json.loads(args.overrides.read_text(encoding="utf-8")) if args.overrides else None
    out = build(args.out, video=args.video, script=args.script, overrides=overrides)
    print(f"{len(apply_overrides(SLIDES, overrides) if overrides else SLIDES)} slides -> {out}")
    if args.video:
        try:
            print(f"  playback: {set_playback(out)}")
        except ImportError:
            print("  playback: set Start automatically / Play full screen in PowerPoint by hand")
    words, minutes = speaking_time(args.script)
    print(f"  notes: {words} words, {minutes:.1f} min at 130 wpm")
    if args.render:
        for sheet in render(out, out.parent / "render"):
            print(f"  {sheet}")
